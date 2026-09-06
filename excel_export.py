"""Excel export for the marking tool - two sheets, plain formatting.

Sheet 1 "Summary": one row per competitor, the 11 automated aspects plus
their total out of 42.

Sheet 2 "Full Marking Sheet": one block per competitor with all 24 official
aspects, the 13 human-marked ones left blank for the Chief Expert to fill in.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scoring import (
    ASPECT_DEFS,
    AUTOMATED,
    AUTOMATED_MAX_TOTAL,
    BEHAVIOUR_MAX_TOTAL,
    JUDGEMENT_MAX_TOTAL,
    MANUAL_OBSERVATION,
    AspectResult,
    SUMMARY_COLUMNS,
    automated_total,
)

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
AMBER_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)
DISCLAIMER = (
    "Draft produced by automated tool. Judgement and behaviour marks pending "
    "Chief Expert review. Not a final score."
)


@dataclass
class CompetitorResult:
    competitor_number: str
    aspects: list[AspectResult]
    error: str | None = None  # e.g. "model file could not be opened"


def _autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_summary_sheet(wb: Workbook, competitor_results: list[CompetitorResult]) -> None:
    ws = wb.active
    ws.title = "Summary"

    headers = ["Competitor"] + SUMMARY_COLUMNS + [f"Total (/{AUTOMATED_MAX_TOTAL})"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for cr in competitor_results:
        by_id = {r.id: r for r in cr.aspects}
        row = [cr.competitor_number]
        for aid in SUMMARY_COLUMNS:
            r = by_id[aid]
            row.append(r.marks_awarded if r.marks_awarded is not None else "")
        row.append(automated_total(cr.aspects))
        ws.append(row)

        row_idx = ws.max_row
        for col_offset, aid in enumerate(SUMMARY_COLUMNS, start=2):
            r = by_id[aid]
            cell = ws.cell(row=row_idx, column=col_offset)
            if r.provisional:
                cell.fill = AMBER_FILL
            elif r.passed is False:
                cell.fill = RED_FILL

    widths = {1: 12}
    for i in range(2, 2 + len(SUMMARY_COLUMNS)):
        widths[i] = 6
    widths[2 + len(SUMMARY_COLUMNS)] = 12
    _autosize(ws, widths)
    ws.freeze_panes = "A2"


FULL_SHEET_HEADERS = [
    "Criterion", "ID", "Aspect", "Type", "Max Mark",
    "Measured Value", "Pass/Fail", "Marks Awarded", "Notes",
]


def _write_full_sheet(wb: Workbook, competitor_results: list[CompetitorResult]) -> None:
    ws = wb.create_sheet("Full Marking Sheet")
    ws.append(["Full Marking Sheet — all 24 aspects, 100 marks"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    for cr in competitor_results:
        ws.append([])  # spacer row
        ws.append([f"Competitor {cr.competitor_number}"])
        title_row = ws.max_row
        ws.cell(row=title_row, column=1).font = Font(bold=True, size=12)

        ws.append(FULL_SHEET_HEADERS)
        header_row = ws.max_row
        for col in range(1, len(FULL_SHEET_HEADERS) + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        for r in cr.aspects:
            type_label = "Automated" if r.kind == AUTOMATED else "Manual"
            pass_fail = "" if r.passed is None else ("PASS" if r.passed else "FAIL")
            ws.append([
                r.criterion, r.id, r.name, type_label, r.max_mark,
                r.measured_value, pass_fail, r.marks_awarded if r.marks_awarded is not None else "",
                r.notes,
            ])
            row_idx = ws.max_row
            if r.provisional:
                for col in range(1, len(FULL_SHEET_HEADERS) + 1):
                    ws.cell(row=row_idx, column=col).fill = AMBER_FILL
            elif r.passed is False:
                for col in range(1, len(FULL_SHEET_HEADERS) + 1):
                    ws.cell(row=row_idx, column=col).fill = RED_FILL

        auto_sub = automated_total(cr.aspects)
        ws.append(["", "", "Automated subtotal", "", "", "", "", f"{auto_sub} / {AUTOMATED_MAX_TOTAL}", ""])
        ws.append(["", "", "Judgement subtotal (Chief Expert)", "", "", "", "", f"_ / {JUDGEMENT_MAX_TOTAL}", ""])
        ws.append(["", "", "Behaviour subtotal (Chief Expert)", "", "", "", "", f"_ / {BEHAVIOUR_MAX_TOTAL}", ""])
        ws.append([DISCLAIMER])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, size=9)

    widths = {1: 10, 2: 6, 3: 26, 4: 11, 5: 10, 6: 34, 7: 10, 8: 15, 9: 50}
    _autosize(ws, widths)


def export_results(
    competitor_results: list[CompetitorResult],
    submissions_folder: Path,
    run_date: date | None = None,
) -> Path:
    run_date = run_date or date.today()
    wb = Workbook()
    _write_summary_sheet(wb, competitor_results)
    _write_full_sheet(wb, competitor_results)

    filename = f"SC2026_Marking_Results_{run_date:%Y-%m-%d}.xlsx"
    output_path = Path(submissions_folder) / filename
    wb.save(output_path)
    return output_path
