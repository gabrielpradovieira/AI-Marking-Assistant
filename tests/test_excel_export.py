import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from excel_export import CompetitorResult, export_results
from scoring import automated_total, score_competitor
from tests.test_scoring import _base_measurements


def test_export_creates_two_sheets_with_expected_rows(tmp_path):
    m_clean = _base_measurements(competitor_number="07")
    m_fail = _base_measurements(
        competitor_number="08",
        psd_width=1920, psd_height=1080, psd_ppi=72.0,
        ngon_count=14, construction_history=True, image_planes=1,
    )
    results = [
        CompetitorResult("07", score_competitor(m_clean)),
        CompetitorResult("08", score_competitor(m_fail)),
    ]

    out_path = export_results(results, tmp_path)
    assert out_path.exists()
    assert out_path.name.startswith("SC2026_Marking_Results_")

    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Summary", "Full Marking Sheet"]

    ws = wb["Summary"]
    header = [c.value for c in ws[1]]
    assert header[0] == "Competitor"
    assert "Total (/42)" in header

    row07 = [c.value for c in ws[2]]
    row08 = [c.value for c in ws[3]]
    assert row07[0] == "07"
    assert row08[0] == "08"
    assert row07[-1] == automated_total(results[0].aspects)
    assert row08[-1] == automated_total(results[1].aspects)

    full = wb["Full Marking Sheet"]
    all_text = "\n".join(
        str(c.value) for row in full.iter_rows() for c in row if c.value is not None
    )
    assert "Competitor 07" in all_text
    assert "Competitor 08" in all_text
    assert "Draft produced by automated tool" in all_text
    assert "JUDGEMENT — CE" in all_text
    assert "OBSERVATION — CE" in all_text


def test_failed_cells_are_red_and_d5_is_amber(tmp_path):
    m_fail = _base_measurements(
        competitor_number="08",
        psd_width=1920, psd_height=1080, psd_ppi=72.0,
    )
    results = [CompetitorResult("08", score_competitor(m_fail))]
    out_path = export_results(results, tmp_path)

    wb = load_workbook(out_path)
    ws = wb["Summary"]
    header = [c.value for c in ws[1]]
    c2_col = header.index("C2") + 1
    d5_col = header.index("D5") + 1
    c2_cell = ws.cell(row=2, column=c2_col)
    d5_cell = ws.cell(row=2, column=d5_col)

    assert c2_cell.fill.start_color.rgb == "00FFC7CE"  # red - failed
    assert d5_cell.fill.start_color.rgb == "00FFE699"  # amber - always provisional
