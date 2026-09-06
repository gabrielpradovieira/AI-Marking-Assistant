import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from config import ToolConfig
from openpyxl import load_workbook
from runner import run_marking
import maya_launcher
from tests.make_mock_submissions import build_mock_submissions

FAKE_MAYAPY = Path(__file__).resolve().parent / "fake_mayapy_for_runner.py"


def _make_config(submissions_folder: Path) -> ToolConfig:
    return ToolConfig(
        submissions_folder=str(submissions_folder),
        mayapy_path=str(sys.executable),
        concept_deadline_date="31/12/2099",
        concept_deadline_time="12:00",
        model_deadline_date="31/12/2099",
        model_deadline_time="16:30",
        triangle_budget=10000,
        canvas_width=3840,
        canvas_height=2160,
        required_ppi=300,
        concept_filename_pattern="ESC2026_TP50_{n}_Digital-Art.psd",
        model_filename_pattern="ESC2026_TP50_{n}_Diving_Helmet.mb",
    )


def _by_id(aspects, aid):
    return next(a for a in aspects if a.id == aid)


def _by_competitor(outcome, number):
    return next(c for c in outcome.competitor_results if c.competitor_number == number)


def test_full_pipeline_scores_all_four_planted_scenarios(tmp_path, monkeypatch):
    submissions = tmp_path / "submissions"
    build_mock_submissions(submissions)
    cfg = _make_config(submissions)

    monkeypatch.setattr(maya_launcher, "worker_script_path", lambda: FAKE_MAYAPY)

    events = []
    outcome = run_marking(cfg, mayapy_path=Path(sys.executable),
                           on_progress=lambda e, d: events.append((e, d)))

    assert len(outcome.competitor_results) == 4
    assert outcome.excel_path is not None and outcome.excel_path.exists()
    assert outcome.log_path is not None and outcome.log_path.exists()
    assert ("phase", "file_checks") in events
    assert ("phase", "scoring") in events
    assert ("phase", "done") in events

    # --- Scenario 1: competitor 07, fully clean pass ---
    c07 = _by_competitor(outcome, "07")
    for r in c07.aspects:
        if r.kind == "automated":
            assert r.passed is True, f"07/{r.id} unexpectedly failed: {r.measured_value}"

    # --- Scenario 2: competitor 08, wrong PPI and canvas size ---
    c08 = _by_competitor(outcome, "08")
    c2 = _by_id(c08.aspects, "C2")
    c3 = _by_id(c08.aspects, "C3")
    assert c2.passed is False and "1920x1080" in c2.measured_value
    assert c3.passed is False and "72" in c3.measured_value
    # Everything else about 08 should still be fine
    assert _by_id(c08.aspects, "A1").passed is True
    assert _by_id(c08.aspects, "D3").passed is True

    # --- Scenario 3: competitor 09, n-gons + history + image plane ---
    c09 = _by_competitor(outcome, "09")
    d3 = _by_id(c09.aspects, "D3")
    d6 = _by_id(c09.aspects, "D6")
    d1 = _by_id(c09.aspects, "D1")
    assert d3.passed is False and "14" in d3.measured_value
    assert d6.passed is False
    assert d1.passed is False and "image plane" in d1.measured_value
    # Its PSD and filenames are otherwise fine
    assert _by_id(c09.aspects, "C2").passed is True
    assert _by_id(c09.aspects, "C3").passed is True

    # --- Scenario 4: competitor 10, missing model file entirely ---
    c10 = _by_competitor(outcome, "10")
    a1 = _by_id(c10.aspects, "A1")
    assert a1.passed is False and "model missing" in a1.measured_value
    for aid in ("D1", "D2", "D3", "D4", "D5", "D6"):
        r = _by_id(c10.aspects, aid)
        assert r.passed is False
        assert r.marks_awarded == 0
    # Its PSD is fine and shouldn't be penalised
    assert _by_id(c10.aspects, "C2").passed is True

    # --- Cross-check against the actual Excel file on disk ---
    wb = load_workbook(outcome.excel_path)
    ws = wb["Summary"]
    header = [c.value for c in ws[1]]
    rows = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    assert set(rows.keys()) == {"07", "08", "09", "10"}
    total_col = header.index("Total (/42)")
    assert rows["07"][total_col].value == 42


def test_maya_unavailable_marks_d_aspects_not_run_without_crashing(tmp_path):
    submissions = tmp_path / "submissions"
    build_mock_submissions(submissions)
    cfg = _make_config(submissions)

    outcome = run_marking(cfg, mayapy_path=None)

    assert outcome.maya_available is False
    c07 = _by_competitor(outcome, "07")
    for aid in ("D1", "D2", "D3", "D4", "D5", "D6"):
        r = _by_id(c07.aspects, aid)
        assert r.passed is False
        assert r.measured_value == "NOT RUN"
    # File/PSD checks should still have run normally
    assert _by_id(c07.aspects, "A1").passed is True
    assert _by_id(c07.aspects, "C2").passed is True
    assert outcome.excel_path is not None and outcome.excel_path.exists()
